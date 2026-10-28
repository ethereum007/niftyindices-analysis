"""Align a user-supplied public-mandate snapshot; never infer paid holdings/popularity."""
import hashlib,json,shutil,sys
from pathlib import Path
import openpyxl
BASE=Path(__file__).parent
NEW=[
 ('itservices','IT, Software & Digital Platforms','IT, software & digital platforms',['IT services','Engineering R&D','Enterprise software','Digital platforms','Cybersecurity'],['Nifty IT','Nifty India Digital']),
 ('banking','Banks & NBFCs','Banks & NBFCs',['Private banks','PSU banks','NBFC lending','Microfinance','Credit infrastructure'],['Nifty Bank','Nifty NBFC']),
 ('utilities','Power & Utilities','Power, utilities & energy',['Power generation','Distribution utilities','Transmission','Power equipment','Energy trading'],['Nifty Power']),
 ('infrastructure','Infrastructure & Construction','Infrastructure & construction',['Roads & highways','Civil EPC','Urban infrastructure','Construction materials','Infrastructure operators'],['Nifty Infrastructure','Nifty Construction']),
 ('logistics','Logistics, Transport & Aviation','Logistics, transport & aviation',['Ports & shipping','Warehousing','Road logistics','Aviation','Express delivery'],['Nifty Transportation & Logistics','Nifty Commercial & Transport Services']),
 ('electronics','Electronics & Semiconductors','Electronics & semiconductors',['EMS & assembly','Semiconductor design','OSAT & packaging','Electronic components','Industrial electronics'],['Nifty India Digital','Nifty Capital Goods']),
 ('insurance','Insurance','Insurance',['Life insurance','General insurance','Health insurance','Reinsurance','Distribution platforms'],['Nifty Insurance']),
 ('waterwaste','Water & Waste Management','Water & waste management',['Water treatment','Wastewater EPC','Pumps & fluid handling','Solid waste services','Resource recovery'],['Nifty Infrastructure']),
 ('chemicals','Chemicals','Chemicals',['Specialty chemicals','Agrochemicals','Fluorochemicals','Commodity chemicals','Industrial gases'],['Nifty Chemicals']),
 ('services','Diversified Services','Diversified services',['Business services','Staffing','Facility services','Professional platforms','Service outsourcing'],['Nifty Services Sector']),
 ('telecommedia','Telecom & Media','Telecom & media',['Telecom operators','Network equipment','Broadcasting','Content & entertainment','Advertising'],['Nifty Telecommunications','Nifty Media']),
 ('travelleisure','Travel, Leisure & Entertainment','Travel, leisure & entertainment',['Hotels & resorts','Travel platforms','Leisure venues','Restaurants','Entertainment services'],['Nifty India Tourism','Nifty Consumer Services'])]
MAP={'Automobiles, EVs & components':'mobility','AI, data centres & digital infrastructure':'ai','Consumer, FMCG & durables':'consumption','Defence & aerospace':'defence','Pharma & healthcare':'healthcare','Financial services & fintech':'fintech','Manufacturing & industrials':'manufacturing','Renewables, clean energy & recycling':'cleanenergy','Capital markets & wealth management':'capitalmarkets','Metals, mining & minerals':'minerals','Railways':'railways','Real estate, housing & REITs':'housing','Oil & gas':'oilgas',**{r[2]:r[0] for r in NEW}}
EXTRA=[('agriculture','Agriculture & Food Supply Chain',['Seeds & crop inputs','Farm equipment','Irrigation','Food processing','Cold chain'],['Nifty Rural']),('textiles','Textiles & Apparel',['Yarn & fabrics','Garment manufacturing','Technical textiles','Home textiles','Export supply chain'],[]),('packaging','Packaging & Paper',['Paper & pulp','Flexible packaging','Rigid packaging','Industrial containers','Sustainable materials'],[])]
def main(source):
 raw=source.read_bytes();sha=hashlib.sha256(raw).hexdigest();folder=BASE/'data/theme-discovery';folder.mkdir(exist_ok=True)
 shutil.copyfile(source,folder/'smallcase-sector-coverage.xlsx')
 w=openpyxl.load_workbook(source,read_only=True,data_only=True)
 summary=list(w['Sector summary'].iter_rows(values_only=True))
 sectors=[{'name':r[1],'products':r[2],'managers':r[3],'sheet_row':i,'family_id':MAP[r[1]]} for i,r in enumerate(summary,1) if len(r)>2 and r[1] in MAP]
 products=[]
 for i,r in enumerate(w['All smallcases'].iter_rows(values_only=True),1):
  if r and isinstance(r[0],(int,float)):
   products.append({'row':i,'name':r[1],'manager':r[2],'coverage_type':r[3],'sectors':r[4].split('; '),'url':r[13],'id':r[14]})
 assert len(products)==545 and len({r['id'] for r in products})==545
 assert len(sectors)==25
 for sector in sectors:
  matched=[r for r in products if sector['name'] in r['sectors']]
  assert len(matched)==sector['products'],sector['name']
  assert len({r['manager'] for r in matched})==sector['managers'],sector['name']
  sector['references']=[{k:r[k] for k in ('name','manager','url','id','row')} for r in matched]
 data={'snapshot_date':'2026-10-28','source_file':'smallcase-sector-coverage.xlsx','sha256':sha,'unique_products':545,'unique_managers':92,'industry_focused_products':139,'sector_categories':sectors,'coverage_types':[{'name':r[1],'products':r[2]} for r in summary if len(r)>2 and isinstance(r[2],(int,float)) and r[1] in ('Diversified / strategy','Sector / industry theme','Asset allocation / target date','Precious metals / bullion','Ethical / ESG screen','Event / ownership / listing theme','Sector / theme rotation','Business group / PSU ownership')],'note':'User-supplied dated public-mandate classification; counts overlap and measure product listings, not subscribers, AUM, holdings, inflows or returns. Strategy and asset-allocation classifications are not equity sectors.'}
 (folder/'alignment.json').write_text(json.dumps(data,indent=2),encoding='utf-8')
 target=BASE.parent/'index-intelligence/themes.json';themes=json.loads(target.read_text(encoding='utf-8'));catalogue=themes['catalogue'];families=catalogue['families'];ids={f['id'] for f in families}
 for id,name,sector,subthemes,related in NEW:
  if id not in ids:families.append({'id':id,'name':name,'description':'Research universe for '+name.lower()+'. Company exposure requires independent primary-source verification.','sources':[],'subthemes':subthemes,'related_indices':related,'membership_status':'Company mapping pending','mapping_note':'Related official indices provide context only, including broad proxies where no dedicated series exists. Their returns and constituents are not custom-theme membership or performance.','company_data_ready':False,'origin':'workbook_alignment'})
 for id,name,subs,related in EXTRA:
  if id not in ids:families.append({'id':id,'name':name,'description':'Additional research proposal beyond the supplied sector categories; companies not yet mapped.','sources':[],'subthemes':subs,'related_indices':related,'membership_status':'Additional research proposal','mapping_note':'Related official indices provide context only. A broad proxy is not a dedicated theme return; absent series remain unavailable.','company_data_ready':False,'origin':'additional_proposal'})
 for sector in sectors:
  f=next(f for f in families if f['id']==sector['family_id']);f['workbook_alignment']={k:sector[k] for k in ('name','products','managers','sheet_row','references')}
  if f.get('origin')=='workbook_alignment':f['sources']=[r['url'] for r in sector['references'][:3]]
 catalogue['workbook_alignment']={k:v for k,v in data.items() if k!='sector_categories'}
 catalogue['workbook_alignment']['mappings']=[{k:s[k] for k in ('name','family_id','products','managers','sheet_row')} for s in sectors]
 catalogue['reviewed_on']='2026-10-28'
 target.write_text(json.dumps(themes,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
 print(json.dumps({'families':len(families),'mapped_workbook_categories':len(sectors),'new_workbook_families':len(NEW),'additional_proposals':len(EXTRA),'source_sha256':sha}))
if __name__=='__main__':main(Path(sys.argv[1]))
